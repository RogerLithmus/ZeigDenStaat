#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
enrich_data.py – Robust data enrichment pipeline for ZeigDenStaat.

Steps:
1. Read existing 'Übersicht', 'Behörden', and 'Unternehmen' sheets from c:\Projects\ZeigDenStaat\data\bundesorganisationen.xlsx.
2. Merge worksheets to build a master record dictionary for all 209 entities.
3. Query Wikipedia OpenSearch API to resolve Wikipedia links for the 79 missing ones.
4. Visit each Wikipedia page and scrape employee count, founding year, budget/etat, website, headquarters, and legal form.
5. Employs a robust double-cell <tr> parser to successfully read company infoboxes (which use <td> fette-Zellen instead of <th>).
6. Clean values (remove Wikipedia citations like [1], clean whitespaces).
7. Apply fuzzy matching against anschriftenverzeichnis.csv to fill in missing address details.
8. Re-generate data/bundesorganisationen.xlsx with exact beautiful openpyxl styling.
"""

import csv
import logging
import os
import re
import sys
import time
import urllib.parse
from difflib import SequenceMatcher
from collections import defaultdict

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ──
BASE_DIR = r"c:\Projects\ZeigDenStaat"
RAW_DIR = os.path.join(BASE_DIR, "data", "raw")
ANSCHRIFT_PATH = os.path.join(RAW_DIR, "anschriftenverzeichnis.csv")
EXCEL_PATH = os.path.join(BASE_DIR, "data", "bundesorganisationen.xlsx")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

HEADERS = {
    "User-Agent": "ZeigDenStaatBot/1.0 (https://github.com/zeigdenstaat; research project) Python-requests"
}
DELAY = 0.5  # Polite delay between Wikipedia requests

# ═══════════════════════════════════════════════════════════════════════════
# Helpers & Cleaning Functions
# ═══════════════════════════════════════════════════════════════════════════

def clean_text(text: str | None) -> str:
    """Strip whitespace, remove zero-width characters, and remove Wikipedia citations."""
    if text is None:
        return ""
    t = str(text)
    # Remove soft hyphens, non-breaking spaces, zero-width spaces
    t = t.replace("\u00ad", "").replace("\u00a0", " ").replace("\u200b", "")
    # Remove Wikipedia reference citations like [1], [2], [Anm. 1], [a]
    t = re.sub(r"\[.*?\]", "", t)
    # Remove multiple spaces/newlines
    t = " ".join(t.split()).strip()
    return t


def clean_search_name(name: str) -> str:
    """Remove parenthetical text for better Wikipedia search matching."""
    name = re.sub(r"\s*\(.*?\)", "", name)
    name = name.strip()
    return name


def normalize_name(name: str) -> str:
    """Normalize entity name for matching."""
    name = name.lower().strip()
    for suffix in [" gmbh", " ag", " ggmbh", " mbh", " se", " e.v.", " ev"]:
        name = name.replace(suffix, "")
    name = re.sub(r"\s*\(.*?\)", "", name)
    name = " ".join(name.split())
    return name


def fuzzy_match(name: str, candidates: list, threshold: float = 0.75) -> str | None:
    """Find best fuzzy match in a list of candidate strings."""
    norm = normalize_name(name)
    best_score = 0
    best_match = None

    for cand in candidates:
        cand_norm = normalize_name(cand)
        score = SequenceMatcher(None, norm, cand_norm).ratio()
        if score > best_score and score >= threshold:
            best_score = score
            best_match = cand

    return best_match


# ═══════════════════════════════════════════════════════════════════════════
# Wikipedia Data Retrieval
# ═══════════════════════════════════════════════════════════════════════════

def find_wikipedia_url(name: str) -> str | None:
    """Query Wikipedia OpenSearch API to find the page URL for a name."""
    cleaned_name = clean_search_name(name)
    url = f"https://de.wikipedia.org/w/api.php?action=opensearch&search={urllib.parse.quote(cleaned_name)}&limit=1&format=json"
    
    try:
        time.sleep(0.1) # Quick delay
        resp = requests.get(url, headers=HEADERS, timeout=10)
        if resp.status_code == 200:
            data = resp.json()
            if len(data) >= 4 and len(data[3]) > 0:
                wiki_url = data[3][0]
                # Avoid standard category or helper pages unless necessary
                if "Kategorie:" not in wiki_url and "Liste_" not in wiki_url:
                    return wiki_url
    except Exception as e:
        log.warning(f"Error searching Wikipedia for '{cleaned_name}': {e}")
    return None


def scrape_wikipedia_infobox(url: str) -> dict:
    """Scrape the infobox table of a German Wikipedia article with custom cell parser."""
    result = {
        "Mitarbeiteranzahl": "",
        "Gründungsjahr": "",
        "Hauptsitz": "",
        "Budget/Etat": "",
        "Website": "",
        "Rechtsform": "",
    }
    
    try:
        time.sleep(DELAY)
        resp = requests.get(url, headers=HEADERS, timeout=15)
        if resp.status_code != 200:
            log.warning(f"Wikipedia returned status {resp.status_code} for {url}")
            return result
        
        soup = BeautifulSoup(resp.content, "html.parser")
        
        # Locate the infobox table
        infobox = None
        for cls in ["infobox", "infobox-body", "toptextcells", "vorlage-infobox-unternehmen"]:
            infobox = soup.find("table", class_=re.compile(cls, re.I))
            if infobox:
                break
        if not infobox:
            infobox = soup.find("table", class_=re.compile(r"float-right|wikitable.*infobox", re.I))
        if not infobox:
            return result
        
        # Parse rows
        rows = infobox.find_all("tr")
        for row in rows:
            # Check for two columns in a row (typical key-value pairing)
            cells = row.find_all(["th", "td"], recursive=False)
            if len(cells) == 2:
                label = clean_text(cells[0].get_text())
                value = clean_text(cells[1].get_text())
                label_lower = label.lower()
                
                # Check for standard fields
                if any(k in label_lower for k in ["mitarbeiter", "beschäftigte", "bedienstete", "personal", "mitarbeiterzahl"]):
                    result["Mitarbeiteranzahl"] = value
                elif any(k in label_lower for k in ["gründung", "erricht", "bestehen seit", "gründungsdatum", "gründungsjahr"]):
                    result["Gründungsjahr"] = value
                elif any(k in label_lower for k in ["sitz", "hauptsitz", "standort", "sitz der gesellschaft"]):
                    result["Hauptsitz"] = value
                elif any(k in label_lower for k in ["etat", "haushalt", "budget", "umsatz", "umsatzentwicklung"]):
                    result["Budget/Etat"] = value
                elif any(k in label_lower for k in ["website", "webpräsenz", "webseite", "homepage"]):
                    a_tag = cells[1].find("a", href=True)
                    if a_tag:
                        href = a_tag["href"]
                        if href.startswith("//"):
                            href = "https:" + href
                        if href.startswith("http"):
                            result["Website"] = href
                        else:
                            result["Website"] = value
                    else:
                        result["Website"] = value
                elif any(k in label_lower for k in ["rechtsform", "organisationsform"]):
                    result["Rechtsform"] = value
                    
    except Exception as e:
        log.warning(f"Error scraping Wikipedia page '{url}': {e}")
        
    return result


# ═══════════════════════════════════════════════════════════════════════════
# Anschriftenverzeichnis Support
# ═══════════════════════════════════════════════════════════════════════════

def load_anschriftenverzeichnis(filepath: str) -> dict:
    """Load addresses from Anschriftenverzeichnis CSV into a dict keyed by raw name."""
    lookup = {}
    if not os.path.exists(filepath):
        log.warning(f"Anschriftenverzeichnis file not found: {filepath}")
        return lookup
        
    try:
        df = pd.read_csv(filepath, sep=";", encoding="utf-8")
        for _, row in df.iterrows():
            org = str(row.get("Organisation", "")).strip()
            if org:
                lookup[org] = {
                    "Adresse": clean_text(row.get("Adresse", "")),
                    "PLZ": clean_text(row.get("PLZ", "")),
                    "Ort": clean_text(row.get("Ort", "")),
                    "Telefon": clean_text(row.get("Telefon", "")),
                    "Email": clean_text(row.get("E-Mail", "")),
                    "Website": clean_text(row.get("Internetadresse", "")),
                }
        log.info(f"Loaded {len(lookup)} address entries from Anschriftenverzeichnis")
    except Exception as e:
        log.warning(f"Could not load Anschriftenverzeichnis: {e}")
    return lookup


# ═══════════════════════════════════════════════════════════════════════════
# Core Excel Generation Functions (matching original style)
# ═══════════════════════════════════════════════════════════════════════════

def write_sheet(ws, data: list[dict], columns: list[str], header_fill):
    """Write list of dictionaries to an openpyxl sheet with professional formatting."""
    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    alt_fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

    # Header row
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data rows
    for row_idx, record in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = record.get(col_name, "")
            # Ensure nan or null is empty string
            if pd.isna(value) or str(value) == "nan":
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = wrap_alignment
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Set row heights
    ws.row_dimensions[1].height = 28

    # Auto-width columns
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        # Scan first 50 rows for width matching
        for row_idx in range(2, min(len(data) + 2, 50)):
            cell_val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, min(len(cell_val), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 55)

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto-filter
    if data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(data) + 1}"


# ═══════════════════════════════════════════════════════════════════════════
# Main Program
# ═══════════════════════════════════════════════════════════════════════════

def main():
    print("=" * 70)
    print("ZEIGDENSTAAT - ROBUST DATA ENRICHMENT PROCESS")
    print("=" * 70)

    # 1. Verify files exist
    if not os.path.exists(EXCEL_PATH):
        log.error(f"Excel database not found at: {EXCEL_PATH}")
        sys.exit(1)

    # 2. Read sheets into DataFrames
    log.info(f"Loading Excel file: {EXCEL_PATH}")
    try:
        xl = pd.ExcelFile(EXCEL_PATH)
        df_all = pd.read_excel(xl, sheet_name="Übersicht")
        df_beh = pd.read_excel(xl, sheet_name="Behörden")
        df_unt = pd.read_excel(xl, sheet_name="Unternehmen")
    except Exception as e:
        log.error(f"Failed to read existing Excel file: {e}")
        sys.exit(1)

    log.info(f"Loaded {len(df_all)} entities from 'Übersicht' sheet.")

    # 3. Create a master dictionary of entities to enrich
    # We load additional contact fields from 'Behörden' sheet if available
    records = []
    contact_fields = {} # keyed by Name
    
    for _, row in df_beh.iterrows():
        name = str(row.get("Name", "")).strip()
        if name:
            contact_fields[name] = {
                "Adresse": str(row.get("Adresse", "")) if pd.notna(row.get("Adresse")) else "",
                "PLZ": str(row.get("PLZ", "")) if pd.notna(row.get("PLZ")) else "",
                "Ort": str(row.get("Ort", "")) if pd.notna(row.get("Ort")) else "",
                "Telefon": str(row.get("Telefon", "")) if pd.notna(row.get("Telefon")) else "",
                "E-Mail": str(row.get("E-Mail", "")) if pd.notna(row.get("E-Mail")) else "",
            }

    for _, row in df_all.iterrows():
        name = str(row.get("Name", "")).strip()
        if not name:
            continue
            
        typ = str(row.get("Typ", ""))
        rec = {
            "Name": name,
            "Abkürzung": str(row.get("Abkürzung", "")) if pd.notna(row.get("Abkürzung")) else "",
            "Typ": typ,
            "Untertyp": str(row.get("Untertyp", "")) if pd.notna(row.get("Untertyp")) else "",
            "Geschäftsbereich": str(row.get("Geschäftsbereich", "")) if pd.notna(row.get("Geschäftsbereich")) else "",
            "Hauptsitz": str(row.get("Hauptsitz", "")) if pd.notna(row.get("Hauptsitz")) else "",
            "Ort": str(row.get("Ort", "")) if pd.notna(row.get("Ort")) else "",
            "Gründungsjahr": str(row.get("Gründungsjahr", "")) if pd.notna(row.get("Gründungsjahr")) else "",
            "Mitarbeiteranzahl": str(row.get("Mitarbeiteranzahl", "")) if pd.notna(row.get("Mitarbeiteranzahl")) else "",
            "Budget/Etat": str(row.get("Budget/Etat", "")) if pd.notna(row.get("Budget/Etat")) else "",
            "Rechtsform": str(row.get("Rechtsform", "")) if pd.notna(row.get("Rechtsform")) else "",
            "Website": str(row.get("Website", "")) if pd.notna(row.get("Website")) else "",
            "Wikipedia": str(row.get("Wikipedia", "")) if pd.notna(row.get("Wikipedia")) else "",
            "Adresse": "",
            "PLZ": "",
            "Telefon": "",
            "E-Mail": "",
        }
        
        # Restore contact fields if Behörde
        if name in contact_fields:
            rec["Adresse"] = contact_fields[name]["Adresse"]
            rec["PLZ"] = contact_fields[name]["PLZ"]
            rec["Ort"] = contact_fields[name]["Ort"] or rec["Ort"]
            rec["Telefon"] = contact_fields[name]["Telefon"]
            rec["E-Mail"] = contact_fields[name]["E-Mail"]
            
        # Clean up any nan strings in loaded fields
        for k, v in rec.items():
            if str(v).lower() in ("nan", "none", ""):
                rec[k] = ""
                
        records.append(rec)

    # 4. Load Anschriftenverzeichnis for address gap-filling
    anschrift_lookup = load_anschriftenverzeichnis(ANSCHRIFT_PATH)
    anschrift_names = list(anschrift_lookup.keys())

    # 5. Pipeline execution: Loop and Enrich!
    log.info("Starting systematic link resolution and Wikipedia infobox scraping...")
    
    total = len(records)
    resolved_count = 0
    scraped_count = 0
    
    for idx, rec in enumerate(records, 1):
        name = rec["Name"]
        wiki_url = rec["Wikipedia"]
        
        # A. Resolve Wikipedia URL if missing
        if not wiki_url:
            log.info(f"[{idx}/{total}] Attempting Wikipedia search: {name}")
            found_url = find_wikipedia_url(name)
            if found_url:
                rec["Wikipedia"] = found_url
                wiki_url = found_url
                resolved_count += 1
                log.info(f"  ✓ Resolved Wikipedia URL: {found_url}")
            else:
                log.info("  ✗ No Wikipedia page found via search.")
        
        # B. Scrape from Wikipedia page if URL is present
        if wiki_url:
            log.info(f"[{idx}/{total}] Scraping Wikipedia infobox: {name} ({wiki_url})")
            info = scrape_wikipedia_infobox(wiki_url)
            
            # Enrich missing columns if scrape returned fields
            updated_fields = []
            
            if info.get("Mitarbeiteranzahl") and not rec["Mitarbeiteranzahl"]:
                rec["Mitarbeiteranzahl"] = info["Mitarbeiteranzahl"]
                updated_fields.append("Mitarbeiter")
                
            if info.get("Gründungsjahr") and not rec["Gründungsjahr"]:
                rec["Gründungsjahr"] = info["Gründungsjahr"]
                updated_fields.append("Gründungsjahr")
                
            if info.get("Hauptsitz") and not rec["Hauptsitz"]:
                rec["Hauptsitz"] = info["Hauptsitz"]
                rec["Ort"] = rec["Ort"] or info["Hauptsitz"]
                updated_fields.append("Sitz")
                
            if info.get("Budget/Etat") and not rec["Budget/Etat"]:
                rec["Budget/Etat"] = info["Budget/Etat"]
                updated_fields.append("Budget")
                
            if info.get("Website") and not rec["Website"]:
                rec["Website"] = info["Website"]
                updated_fields.append("Website")
                
            if info.get("Rechtsform") and not rec["Rechtsform"]:
                rec["Rechtsform"] = info["Rechtsform"]
                updated_fields.append("Rechtsform")
                
            if updated_fields:
                scraped_count += 1
                log.info(f"  ✓ Enriched fields: {', '.join(updated_fields)}")
            else:
                log.info("  - No new fields added from Wikipedia infobox.")
        
        # C. Match against Anschriftenverzeichnis to fill remaining address details
        # (especially useful if Sitz, PLZ, Adresse, Telefon or E-Mail is blank)
        if not rec["Adresse"] or not rec["PLZ"] or not rec["Telefon"] or not rec["Website"]:
            matched_name = fuzzy_match(name, anschrift_names, threshold=0.85)
            if matched_name:
                ad = anschrift_lookup[matched_name]
                if not rec["Adresse"]:
                    rec["Adresse"] = ad["Adresse"]
                if not rec["PLZ"]:
                    rec["PLZ"] = ad["PLZ"]
                if not rec["Ort"]:
                    rec["Ort"] = ad["Ort"]
                if not rec["Hauptsitz"]:
                    rec["Hauptsitz"] = ad["Ort"]
                if not rec["Telefon"]:
                    rec["Telefon"] = ad["Telefon"]
                if not rec["E-Mail"]:
                    rec["E-Mail"] = ad["Email"]
                if not rec["Website"]:
                    rec["Website"] = ad["Website"]
                log.info(f"  ✓ Enriched contact details from Anschriftenverzeichnis via match: '{matched_name}'")

        # Small delay between entities (even if no request was made to keep loop smooth)
        time.sleep(0.05)

    # 6. Re-generate beautiful Excel workbook
    log.info("Generating beautifully styled Excel workbook...")
    wb = Workbook()

    # Sheet Colors
    header_fill_all = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid") # Dark Navy
    header_fill_beh = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid") # Slate Blue
    header_fill_unt = PatternFill(start_color="048A81", end_color="048A81", fill_type="solid") # Teal

    # Sheet 1: Übersicht
    ws_all = wb.active
    ws_all.title = "Übersicht"
    all_columns = ["Name", "Abkürzung", "Typ", "Untertyp", "Geschäftsbereich",
                    "Hauptsitz", "Ort", "Gründungsjahr", "Mitarbeiteranzahl",
                    "Budget/Etat", "Rechtsform", "Website", "Wikipedia"]
    write_sheet(ws_all, records, all_columns, header_fill_all)

    # Sheet 2: Behörden
    behoerden = [r for r in records if r["Typ"] == "Behörde"]
    ws_beh = wb.create_sheet("Behörden")
    beh_columns = ["Name", "Abkürzung", "Untertyp", "Geschäftsbereich",
                    "Hauptsitz", "Adresse", "PLZ", "Ort",
                    "Mitarbeiteranzahl", "Budget/Etat",
                    "Telefon", "E-Mail", "Website", "Wikipedia"]
    write_sheet(ws_beh, behoerden, beh_columns, header_fill_beh)

    # Sheet 3: Unternehmen
    unternehmen = [r for r in records if r["Typ"] == "Unternehmen"]
    ws_unt = wb.create_sheet("Unternehmen")
    unt_columns = ["Name", "Rechtsform", "Geschäftsbereich",
                    "Hauptsitz", "Ort",
                    "Mitarbeiteranzahl", "Gründungsjahr", "Budget/Etat",
                    "Website", "Wikipedia"]
    write_sheet(ws_unt, unternehmen, unt_columns, header_fill_unt)

    # Save Excel
    wb.save(EXCEL_PATH)
    log.info(f"Excel file successfully saved: {EXCEL_PATH}")

    # 7. Quality Report / Gap Analysis after Run
    print("\n" + "=" * 70)
    print("ENRICHMENT RESULTS REPORT")
    print("=" * 70)
    print(f"Total processed entities:  {total}")
    print(f"Wikipedia URLs resolved:    {resolved_count}")
    print(f"Entities enriched from Wiki: {scraped_count}")
    
    # Calculate Gaps
    df_new_all = pd.read_excel(EXCEL_PATH, sheet_name="Übersicht")
    print("\nMissing Fields (Null/Empty) in New Übersicht Sheet:")
    nulls = df_new_all.isnull().sum()
    for col, null_cnt in nulls.items():
        filled_cnt = total - null_cnt
        fill_rate = (filled_cnt / total) * 100
        print(f"  - {col:18s} : {null_cnt:3d} missing ({fill_rate:5.1f}% coverage)")
        
    print("\nSpot Checks:")
    bwi_spot = df_new_all[df_new_all["Name"].str.contains("BWI", case=False, na=False)]
    if not bwi_spot.empty:
        print(f"  - BWI GmbH  -> Mitarbeiter: {bwi_spot.iloc[0]['Mitarbeiteranzahl']}, Budget: {bwi_spot.iloc[0]['Budget/Etat']}, Wikipedia: {bwi_spot.iloc[0]['Wikipedia']}")
        
    rki_spot = df_new_all[df_new_all["Name"].str.contains("Robert Koch", case=False, na=False)]
    if not rki_spot.empty:
        print(f"  - RKI       -> Mitarbeiter: {rki_spot.iloc[0]['Mitarbeiteranzahl']}, Budget: {rki_spot.iloc[0]['Budget/Etat']}")
        
    dest_spot = df_new_all[df_new_all["Name"].str.contains("Statistisches Bundesamt", case=False, na=False)]
    if not dest_spot.empty:
        print(f"  - Destatis  -> Mitarbeiter: {dest_spot.iloc[0]['Mitarbeiteranzahl']}, Budget: {dest_spot.iloc[0]['Budget/Etat']}, Gründung: {dest_spot.iloc[0]['Gründungsjahr']}")
        
    wpk_spot = df_new_all[df_new_all["Name"].str.contains("Wirtschaftsprüferkammer", case=False, na=False)]
    if not wpk_spot.empty:
        print(f"  - WPK       -> Wiki-URL: {wpk_spot.iloc[0]['Wikipedia']}, Ort: {wpk_spot.iloc[0]['Ort']}")
        
    print("=" * 70)


if __name__ == "__main__":
    main()
